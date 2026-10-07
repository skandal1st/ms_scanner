"""Shared report values for desktop and terminal physical counts."""
import io

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS = ['Бренд', 'Товар', 'GTIN', 'Остаток МС', 'В отгрузках выбранных статусов',
           'Ожидалось', 'Посчитано', 'Разница', 'Проверен']


def report_rows(session, counts):
    settings = session.settings or {}
    manual = settings.get('count_method') == 'quantity'
    for row in session.plan:
        entered = row['key'] in counts
        counted = counts.get(row['key'], 0)
        reviewed = ((session.mode == 'acceptance' and session.status == 'completed')
                    or row['folder_name'] in settings.get('reviewed_brands', []))
        if manual:
            reviewed = entered
        yield [row['folder_name'], row.get('product_name') or '',
               row.get('gtin') or ', '.join(row.get('gtins', [])),
               row['base_qty'], row['shipment_qty'], row['expected_qty'],
               counted if not manual or entered else None,
               counted - row['expected_qty'] if reviewed else None,
               'Да' if reviewed else 'Нет']


def build_count_xlsx(session, counts):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Инвентаризация' if session.mode == 'inventory' else 'Сверка приёмки'
    settings = session.settings or {}
    def append(values):
        sheet.append([ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value, str) else value
                      for value in values])
    append([session.name])
    append(['Склад', settings.get('store_name', '')])
    append(['Снимок остатков', settings.get('snapshot_at', '')])
    append(['Статус', {'active': 'В работе', 'completed': 'Завершена'}.get(session.status, session.status)])
    append(['Разница рассчитана только для проверенных позиций. Остатки МойСклада не изменяются.'])
    append(HEADERS)
    for row in report_rows(session, counts):
        append(row)

    # Product names and GTINs are data, including strings beginning with '='.
    for cells in sheet:
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    sheet['A1'].font = Font(size=16, bold=True, color='174EA6')
    sheet.row_dimensions[1].height = 26
    for row_number in (1, 5):
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=9)
    sheet.row_dimensions[5].height = 32
    sheet.row_dimensions[6].height = 44
    for cell in sheet[6]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='174EA6')
    for row in sheet.iter_rows(min_row=7):
        row[2].number_format = '@'
        for cell in row[3:8]:
            cell.number_format = '0.###;[Red]-0.###;0'
        if row[7].value is not None and row[7].value != 0:
            row[7].fill = PatternFill('solid', fgColor='FDE9E7')
    for index, width in enumerate((24, 48, 30, 16, 23, 16, 16, 16, 14), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'D7'
    sheet.auto_filter.ref = f'A6:I{sheet.max_row}'
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
