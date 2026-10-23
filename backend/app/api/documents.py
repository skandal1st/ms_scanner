from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func
from pydantic import BaseModel
from typing import Optional, List
from uuid import UUID
from datetime import datetime
import httpx

from app.db.session import get_db
from app.db.models import User, Document, DocumentKind, DocumentStatus, Integration, OrganizationProfile, Scan
from app.api.deps import get_current_user, get_active_organization_profile
from app.services.moysklad import (MoySkladService, SUPPORTED_KINDS, customer_order_links,
    customer_order_empty_message, customer_order_direct_shipment_count, shipment_matches_customer_order)
from app.services.customer_order_filters import resolve_order_filter
from app.core.security import decrypt_token
from app.services.collection_start import collection_started, require_collection_started
from app.services.shipment_labels import update_shipment_metadata, shipment_display_name
from app.services.shipment_sessions import shipment_session_query, visible_documents
from app.schemas.upd_export import UpdExportOptions

router = APIRouter(prefix="/documents", tags=["documents"])


class PlanItem(BaseModel):
    gtin: Optional[str]
    gtins: List[str] = []
    pack_gtins: List[str] = []
    pack_quantities: dict[str, int] = {}
    product_id: Optional[str]
    product_name: str
    folder_id: Optional[str] = None
    folder_name: Optional[str] = None
    expected_qty: int
    marked: Optional[bool] = None


class DocumentResponse(BaseModel):
    display_name: Optional[str] = None
    collection_started: bool = False
    id: UUID
    moysklad_id: Optional[str]
    name: str
    kind: DocumentKind
    status: DocumentStatus
    scan_count: int = 0
    plan: List[PlanItem] = []
    error_message: Optional[str] = None
    processing_progress: Optional[dict] = None
    cz_doc_ids: Optional[list[dict]] = None
    organization_profile_id: Optional[UUID] = None
    moysklad_organization_id: Optional[str] = None
    moysklad_store_id: Optional[str] = None
    customer_order_name: Optional[str] = None
    created_at: datetime

    model_config = {"from_attributes": True}


class CreateDocumentRequest(BaseModel):
    name: str
    kind: DocumentKind = DocumentKind.demand
    moysklad_id: Optional[str] = None


class MoySkladDocumentItem(BaseModel):
    state_name: Optional[str] = None
    state_color: Optional[int] = None
    id: str
    name: str
    moment: Optional[str]
    customer_order_name: Optional[str] = None
    agent_name: Optional[str] = None
    shipment_count: Optional[int] = None
    retail_sale_count: int = 0
    empty_shipments_message: Optional[str] = None


def _doc_to_response(doc: Document, scan_count: int = 0) -> DocumentResponse:
    # `doc.scans` НЕЛЬЗЯ дёргать из sync-кода — это lazy-relationship,
    # а сессия async (asyncpg). MissingGreenlet → 500.
    # Вызывающий должен передать scan_count явно (либо посчитать SELECT count,
    # либо использовать selectinload, либо знать что документ только что создан).
    return DocumentResponse(
        collection_started=collection_started(doc),
        id=doc.id,
        moysklad_id=doc.moysklad_id,
        name=doc.name,
        display_name=shipment_display_name(doc) if doc.kind == DocumentKind.demand else doc.name,
        kind=doc.kind,
        status=doc.status,
        scan_count=scan_count,
        plan=doc.plan or [],
        error_message=doc.error_message,
        processing_progress=doc.processing_progress,
        cz_doc_ids=doc.cz_doc_ids,
        organization_profile_id=doc.organization_profile_id,
        moysklad_organization_id=doc.moysklad_organization_id,
        moysklad_store_id=doc.moysklad_store_id,
        customer_order_name=getattr(doc, "customer_order_name", None),
        created_at=doc.created_at,
    )


async def _scan_count(db: AsyncSession, document_id: UUID) -> int:
    result = await db.execute(
        select(func.count(Scan.id)).where(Scan.document_id == document_id)
    )
    return int(result.scalar() or 0)


async def _scan_counts(db: AsyncSession, document_ids: List[UUID]) -> dict[UUID, int]:
    if not document_ids:
        return {}
    result = await db.execute(
        select(Scan.document_id, func.count(Scan.id))
        .where(Scan.document_id.in_(document_ids))
        .group_by(Scan.document_id)
    )
    return {doc_id: int(cnt) for doc_id, cnt in result.all()}


async def _get_ms_service(
    current_user: User, db: AsyncSession
) -> MoySkladService:
    result = await db.execute(
        select(Integration).where(Integration.user_id == current_user.id)
    )
    integration = result.scalar_one_or_none()
    if not integration or not integration.moysklad_token:
        raise HTTPException(status_code=400, detail="МойСклад не подключён")
    token = decrypt_token(integration.moysklad_token)
    return MoySkladService(token)


def _ensure_supported_kind(kind: str) -> None:
    if kind not in SUPPORTED_KINDS:
        raise HTTPException(
            status_code=400,
            detail=(
                "Поддерживаются только отгрузочные документы: "
                f"{', '.join(sorted(SUPPORTED_KINDS))}"
            ),
        )


def _customer_order_error(exc: Exception) -> HTTPException:
    if isinstance(exc, httpx.HTTPStatusError) and exc.response.status_code == 403:
        return HTTPException(403, "Нет доступа к заказам покупателей. Обновите права решения в МойСкладе.")
    return HTTPException(502, "Не удалось загрузить заказ или его отгрузки из МойСклада. Повторите попытку.")


@router.get("/customer-orders", response_model=List[MoySkladDocumentItem])
async def list_customer_orders(
    search: Optional[str] = None, offset: int = Query(0, ge=0),
    current_user: User = Depends(get_current_user),
    profile: OrganizationProfile = Depends(get_active_organization_profile), db: AsyncSession = Depends(get_db),
    filter_id: Optional[UUID] = None, cursor: Optional[int] = None, response: Response = None,
):
    if cursor is not None and cursor < 0:
        raise HTTPException(400, "Некорректная страница списка заказов")
    selected_filter = resolve_order_filter(profile, filter_id)
    ms = await _get_ms_service(current_user, db)
    try:
        conditions = {"order_filter": selected_filter} if selected_filter else {}
        if cursor is not None:
            rows, next_cursor = await ms.get_customer_order_page(profile.moysklad_organization_id, search,
                                                                 cursor=cursor, **conditions)
            response.headers["X-Next-Offset"] = str(next_cursor) if next_cursor is not None else ""
        else:
            rows = await ms.get_customer_orders(profile.moysklad_organization_id, search, offset=offset, **conditions)
    except Exception as exc:
        raise _customer_order_error(exc) from exc
    return [MoySkladDocumentItem(id=row["id"], name=row.get("name") or "Без номера",
        moment=row.get("moment"), agent_name=(row.get("agent") or {}).get("name"),
        shipment_count=customer_order_direct_shipment_count(row),
        state_name=(row.get('state') or {}).get('name'), state_color=(row.get('state') or {}).get('color'),
        retail_sale_count=len(customer_order_links(row, "retaildemand")),
        empty_shipments_message=customer_order_empty_message(row)) for row in rows]


@router.get("/customer-orders/{order_id}/shipments", response_model=List[MoySkladDocumentItem])
async def customer_order_shipments(
    order_id: UUID, current_user: User = Depends(get_current_user),
    profile: OrganizationProfile = Depends(get_active_organization_profile), db: AsyncSession = Depends(get_db),
):
    ms = await _get_ms_service(current_user, db)
    try:
        order = await ms.get_customer_order(str(order_id))
        if profile.moysklad_organization_id and _ref_id(order, "organization") != profile.moysklad_organization_id:
            raise HTTPException(403, "Заказ покупателя относится к другому юрлицу")
        rows = await ms.get_customer_order_demands(str(order_id), profile.moysklad_organization_id, order=order)
    except HTTPException:
        raise
    except Exception as exc:
        raise _customer_order_error(exc) from exc
    completed = set((await db.execute(select(Document.moysklad_id).where(
        Document.user_id == current_user.id, Document.organization_profile_id == profile.id,
        Document.kind == DocumentKind.demand, Document.status == DocumentStatus.accepted,
        Document.moysklad_id.in_([row["id"] for row in rows]),
    ))).scalars().all())
    return [MoySkladDocumentItem(id=row["id"], name=row.get("name") or "Без номера", moment=row.get("moment"),
        customer_order_name=order.get("name"), agent_name=(row.get("agent") or {}).get("name")) for row in rows
        if row["id"] not in completed and not row.get("deleted")
        and (not profile.moysklad_organization_id or _ref_id(row, "organization") == profile.moysklad_organization_id)]


@router.get("/moysklad/{kind}", response_model=List[MoySkladDocumentItem])
async def list_moysklad_documents(
    kind: str,
    search: Optional[str] = Query(None, description="Поиск по номеру/контрагенту (полнотекстовый поиск МС)"),
    current_user: User = Depends(get_current_user),
    profile: OrganizationProfile = Depends(get_active_organization_profile),
    db: AsyncSession = Depends(get_db),
):
    """Список МС-документов выбранного типа (supply/demand/loss/move/salesreturn)."""
    _ensure_supported_kind(kind)
    ms = await _get_ms_service(current_user, db)
    return await ms.get_documents(
        kind,
        search=search,
        organization_id=profile.moysklad_organization_id,
    )


@router.get("/moysklad-supplies", response_model=List[MoySkladDocumentItem])
async def list_moysklad_supplies_alias(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Старый алиас приёмки больше не используется."""
    raise HTTPException(status_code=410, detail="Приёмка отключена. Используйте отгрузочные документы.")


@router.get("/", response_model=List[DocumentResponse])
async def list_documents(
    kind: Optional[DocumentKind] = None,
    current_user: User = Depends(get_current_user),
    profile: OrganizationProfile = Depends(get_active_organization_profile),
    db: AsyncSession = Depends(get_db),
):
    query = select(Document).where(
        Document.user_id == current_user.id,
        Document.organization_profile_id == profile.id,
    )
    if kind is not None:
        _ensure_supported_kind(kind.value)
        query = query.where(Document.kind == kind)
    query = query.order_by(Document.created_at.desc())
    result = await db.execute(query)
    documents = result.scalars().all()
    counts = await _scan_counts(db, [d.id for d in documents])
    return [_doc_to_response(d, counts.get(d.id, 0)) for d in visible_documents(documents, counts)]


def _plan_source_kind(doc_kind: str) -> str:
    """Тип МС-документа, из позиций которого строится план.

    Списание (loss) не имеет собственных позиций в МС — его план берём из
    привязанной отгрузки (demand): «прикрепляем отгрузку как план списания».
    Для остальных типов источник плана совпадает с типом документа.
    """
    return "demand" if doc_kind == "loss" else doc_kind


def _ref_id(payload: dict, field: str) -> Optional[str]:
    ref = payload.get(field) or {}
    meta = ref.get("meta") or {}
    if ref.get("id"):
        return str(ref["id"])
    href = meta.get("href") or ""
    return href.rstrip("/").rsplit("/", 1)[-1] if href else None


async def _profile_for_organization(
    db: AsyncSession,
    user_id,
    organization_id: Optional[str],
    fallback: OrganizationProfile,
) -> OrganizationProfile:
    if not organization_id or fallback.moysklad_organization_id == organization_id:
        return fallback
    found = (
        await db.execute(
            select(OrganizationProfile).where(
                OrganizationProfile.user_id == user_id,
                OrganizationProfile.moysklad_organization_id == organization_id,
            )
        )
    ).scalar_one_or_none()
    if found:
        return found
    if fallback.moysklad_organization_id is None:
        fallback.moysklad_organization_id = organization_id
        return fallback
    profile = OrganizationProfile(
        user_id=user_id,
        moysklad_organization_id=organization_id,
        name=f"Юрлицо {organization_id[:8]}",
        is_default=False,
    )
    db.add(profile)
    await db.flush()
    return profile


@router.post("/", response_model=DocumentResponse, status_code=201)
async def create_document(
    body: CreateDocumentRequest,
    current_user: User = Depends(get_current_user),
    profile: OrganizationProfile = Depends(get_active_organization_profile),
    db: AsyncSession = Depends(get_db),
):
    _ensure_supported_kind(body.kind.value)
    verified_ms_doc = None
    if body.moysklad_id and body.kind == DocumentKind.demand:
        from app.services.document_guard import lock_ms_document
        await lock_ms_document(db, current_user.id, body.kind, body.moysklad_id)
        ms = await _get_ms_service(current_user, db)
        try:
            verified_ms_doc = await ms.get_document('demand', body.moysklad_id)
        except Exception as exc:
            raise HTTPException(502, 'Не удалось проверить отгрузку в МойСкладе. Повторите попытку.') from exc
        from app.services.shipment_guard import ensure_active_shipment, ensure_no_stranded_scans
        try:
            ensure_active_shipment(verified_ms_doc)
        except ValueError as exc:
            raise HTTPException(409, str(exc)) from exc
        existing = (await db.execute(shipment_session_query(current_user.id, body.moysklad_id))).scalars().first()
        if existing is not None:
            if existing.status == DocumentStatus.accepted:
                return _doc_to_response(existing, await _scan_count(db, existing.id))
            await ensure_no_stranded_scans(db, ms, current_user.id, existing.organization_profile_id,
                                          existing.moysklad_customer_order_id, body.moysklad_id)
            update_shipment_metadata(existing, verified_ms_doc)
            await db.commit()
            return _doc_to_response(existing, await _scan_count(db, existing.id))
    plan: list = []
    ms_doc = verified_ms_doc or {}
    ms_organization_id: Optional[str] = None
    ms_store_id: Optional[str] = None
    if body.moysklad_id:
        # Подгружаем план сборки из МС-документа: positions → expected_qty по товарам.
        # Если МС не подключён или запрос упал — план остаётся пустым (произвольная сборка).
        try:
            ms = await _get_ms_service(current_user, db)
            ms_doc = verified_ms_doc or await ms.get_document(_plan_source_kind(body.kind.value), body.moysklad_id)
            ms_organization_id = _ref_id(ms_doc, "organization")
            ms_store_id = _ref_id(ms_doc, "store")
            profile = await _profile_for_organization(
                db, current_user.id, ms_organization_id, profile
            )
            plan = await ms.build_plan(_plan_source_kind(body.kind.value), body.moysklad_id)
        except HTTPException:
            pass
        except Exception as exc:
            # build_plan может упасть на любом этапе; не блокируем создание документа
            from app.core.logging import logger as _lg
            _lg.warning(
                "create_document.build_plan_failed",
                kind=body.kind.value,
                moysklad_id=body.moysklad_id,
                error=str(exc),
            )

    doc = Document(
        user_id=current_user.id,
        name=body.name,
        kind=body.kind,
        moysklad_id=body.moysklad_id,
        organization_profile_id=profile.id,
        moysklad_organization_id=ms_organization_id,
        moysklad_store_id=ms_store_id,
        plan=plan,
    )
    db.add(doc)
    if body.kind == DocumentKind.demand and ms_doc:
        update_shipment_metadata(doc, ms_doc)
    await db.commit()
    await db.refresh(doc)
    # Документ только что создан — сканов заведомо нет.
    return _doc_to_response(doc, scan_count=0)


class ResolveDocRequest(BaseModel):
    moysklad_id: str
    kind: DocumentKind = DocumentKind.demand
    customer_order_id: Optional[UUID] = None


@router.post("/resolve", response_model=DocumentResponse)
async def resolve_document(
    body: ResolveDocRequest,
    current_user: User = Depends(get_current_user),
    profile: OrganizationProfile = Depends(get_active_organization_profile),
    db: AsyncSession = Depends(get_db),
):
    """Найти-или-создать наш Document по документу МС (kind + moysklad_id).

    Повторное открытие отгрузки возвращает сохранённую сборку, включая завершённую.
    Из заказа завершённую сборку нельзя запускать повторно."""
    _ensure_supported_kind(body.kind.value)

    order_name = None
    from app.services.document_guard import lock_ms_document
    await lock_ms_document(db, current_user.id, body.kind, body.moysklad_id)
    verified_ms_doc = None
    if body.customer_order_id:
        if body.kind != DocumentKind.demand:
            raise HTTPException(400, "Из заказа покупателя можно открыть только отгрузку")
        ms = await _get_ms_service(current_user, db)
        try:
            verified_ms_doc = await ms.get_document("demand", body.moysklad_id)
            order = await ms.get_customer_order(str(body.customer_order_id))
        except Exception as exc:
            raise _customer_order_error(exc) from exc
        if not shipment_matches_customer_order(verified_ms_doc, order, str(body.customer_order_id)):
            raise HTTPException(409, "Отгрузка больше не связана с выбранным заказом. Обновите список отгрузок.")
        if profile.moysklad_organization_id and any(_ref_id(item, "organization") != profile.moysklad_organization_id
                                                 for item in (order, verified_ms_doc)):
            raise HTTPException(403, "Заказ или отгрузка относятся к другому юрлицу")
        order_name = order.get("name")

    from app.services.shipment_guard import ensure_active_shipment, ensure_no_stranded_scans
    if body.kind == DocumentKind.demand:
        if verified_ms_doc is None:
            ms = await _get_ms_service(current_user, db)
            try:
                verified_ms_doc = await ms.get_document('demand', body.moysklad_id)
            except Exception as exc:
                raise HTTPException(502, 'Не удалось проверить отгрузку в МойСкладе. Повторите попытку.') from exc
        try:
            ensure_active_shipment(verified_ms_doc)
        except ValueError as exc:
            raise HTTPException(409, str(exc)) from exc

    existing_query = (shipment_session_query(current_user.id, body.moysklad_id,
                      profile.id if body.customer_order_id else None)
                      if body.kind == DocumentKind.demand else
                      select(Document).where(Document.user_id == current_user.id,
                          Document.moysklad_id == body.moysklad_id, Document.kind == body.kind,
                          Document.status != DocumentStatus.accepted).order_by(Document.created_at.desc()))
    existing = (
        (
            await db.execute(
                existing_query
            )
        )
        .scalars()
        .first()
    )
    if existing is not None:
        if body.customer_order_id and existing.status == DocumentStatus.accepted:
            raise HTTPException(409, 'Отгрузка уже собрана. Обновите список отгрузок заказа.')
        if body.kind == DocumentKind.demand and existing.status == DocumentStatus.accepted:
            return _doc_to_response(existing, await _scan_count(db, existing.id))
        if body.kind == DocumentKind.demand:
            await ensure_no_stranded_scans(db, ms, current_user.id, existing.organization_profile_id,
                                          body.customer_order_id or existing.moysklad_customer_order_id,
                                          body.moysklad_id)
        from app.services.legacy_pack_plan import refresh_legacy_pack_plan
        if existing.status == DocumentStatus.draft and any('pack_quantities' not in p for p in existing.plan or []):
            await refresh_legacy_pack_plan(db, existing, await _get_ms_service(current_user, db), body.kind.value)
        if body.customer_order_id:
            existing.moysklad_customer_order_id = str(body.customer_order_id)
            existing.customer_order_name = order_name
        if body.kind == DocumentKind.demand:
            update_shipment_metadata(existing, verified_ms_doc, order_name)
        if body.customer_order_id or body.kind == DocumentKind.demand:
            await db.commit()
        return _doc_to_response(existing, await _scan_count(db, existing.id))

    ms = await _get_ms_service(current_user, db)
    name: Optional[str] = None
    ms_doc: dict = {}
    try:
        ms_doc = verified_ms_doc or await ms.get_document(body.kind.value, body.moysklad_id)
        name = ms_doc.get("name")
    except Exception:
        pass  # имя не критично — подставим дефолт
    ms_organization_id = _ref_id(ms_doc, "organization")
    ms_store_id = _ref_id(ms_doc, "store")
    profile = await _profile_for_organization(
        db, current_user.id, ms_organization_id, profile
    )
    if body.kind == DocumentKind.demand:
        await ensure_no_stranded_scans(db, ms, current_user.id, profile.id,
                                      body.customer_order_id, body.moysklad_id)
    plan: list = []
    try:
        # objectId из кнопки МС — документ самого этого типа, поэтому план строим
        # напрямую по kind (в отличие от create_document, где loss сеется из demand).
        plan = await ms.build_plan(body.kind.value, body.moysklad_id)
    except Exception as exc:
        if body.customer_order_id:
            raise HTTPException(502, "Не удалось загрузить план отгрузки из МойСклада. Повторите попытку.") from exc
        from app.core.logging import logger as _lg
        _lg.warning(
            "resolve_document.build_plan_failed",
            kind=body.kind.value,
            moysklad_id=body.moysklad_id,
            error=str(exc),
        )

    doc = Document(
        user_id=current_user.id,
        name=name or f"Документ {body.moysklad_id[:8]}",
        kind=body.kind,
        moysklad_id=body.moysklad_id,
        organization_profile_id=profile.id,
        moysklad_organization_id=ms_organization_id,
        moysklad_store_id=ms_store_id,
        moysklad_customer_order_id=str(body.customer_order_id) if body.customer_order_id else None,
        customer_order_name=order_name,
        plan=plan,
    )
    db.add(doc)
    if body.kind == DocumentKind.demand:
        update_shipment_metadata(doc, ms_doc, order_name)
    await db.commit()
    await db.refresh(doc)
    return _doc_to_response(doc, scan_count=0)


@router.post("/{document_id}/refresh-plan", response_model=DocumentResponse)
async def refresh_plan(
    document_id: UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Перетянуть план сборки из МС-документа (если он привязан)."""
    from app.services.document_guard import editable_document
    doc = await editable_document(db, document_id, current_user.id)

    _ensure_supported_kind(doc.kind.value)
    if not doc.moysklad_id:
        raise HTTPException(400, "Документ не привязан к МойСклад")
    ms = await _get_ms_service(current_user, db)
    doc.plan = await ms.build_plan(_plan_source_kind(doc.kind.value), doc.moysklad_id)
    await db.commit()
    await db.refresh(doc)
    return _doc_to_response(doc, await _scan_count(db, doc.id))


@router.get("/{document_id}", response_model=DocumentResponse)
async def get_document(
    document_id: UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Document).where(
            Document.id == document_id,
            Document.user_id == current_user.id,
        )
    )
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    _ensure_supported_kind(doc.kind.value)
    from app.services.legacy_pack_plan import refresh_legacy_pack_plan
    if doc.moysklad_id and doc.status == DocumentStatus.draft and any('pack_quantities' not in p for p in doc.plan or []):
        await refresh_legacy_pack_plan(db, doc, await _get_ms_service(current_user, db), _plan_source_kind(doc.kind.value))
    if doc.kind == DocumentKind.demand and doc.moysklad_id and not doc.moysklad_name:
        # One-time enrichment of shipments created before structured display metadata.
        try:
            ms = await _get_ms_service(current_user, db)
            shipment = await ms.get_document('demand', doc.moysklad_id)
            update_shipment_metadata(doc, shipment)
            await db.commit()
        except (HTTPException, httpx.HTTPError):
            # Keep old collections readable if MS is temporarily unavailable.
            from app.core.logging import logger
            logger.warning('shipment.label_refresh_failed', document_id=str(doc.id))
    return _doc_to_response(doc, await _scan_count(db, doc.id))


async def _upd_source(document_id: UUID, user: User, db: AsyncSession, *, include_positions=True):
    from app.services.upd_export import UpdExportError
    doc = await db.scalar(select(Document).where(Document.id == document_id, Document.user_id == user.id))
    if not doc:
        raise HTTPException(404, "Документ не найден")
    if doc.kind != DocumentKind.demand or not doc.moysklad_id:
        raise HTTPException(400, "УПД можно сформировать только для отгрузки МойСклад")
    if doc.status != DocumentStatus.accepted:
        raise HTTPException(409, "Сначала завершите запись марок в МойСклад")
    try:
        source = await (await _get_ms_service(user, db)).get_upd_source(doc.moysklad_id, include_positions=include_positions)
    except UpdExportError as exc:
        raise HTTPException(409, str(exc)) from None
    except httpx.HTTPError:
        raise HTTPException(502, "Не удалось загрузить реквизиты отгрузки из МойСклад") from None
    organization = source.get("organization") or {}
    if doc.moysklad_organization_id and organization.get("id") != doc.moysklad_organization_id:
        raise HTTPException(409, "Юрлицо отгрузки изменилось. Откройте отгрузку заново")
    return source


@router.get("/{document_id}/upd-context")
async def upd_export_context(document_id: UUID, current_user: User = Depends(get_current_user),
                             db: AsyncSession = Depends(get_db), response: Response = None):
    from app.services.upd_export import party_from_ms
    source = await _upd_source(document_id, current_user, db, include_positions=False)
    if response is not None:
        response.headers["Cache-Control"] = "no-store"
    seller = party_from_ms(source.get("organization") or {})
    buyer = party_from_ms(source.get("agent") or {})
    invoice = source.get("invoice") or {}
    warnings = []
    if source.get("buyer_unavailable"):
        warnings.append("Нет доступа к реквизитам покупателя. Обновите дескриптор и переустановите решение МойСклад для получения новых прав.")
    elif (not buyer.inn or not buyer.address or (len(buyer.inn) == 10 and not buyer.kpp)
          or (len(buyer.inn) == 12 and (not buyer.first_name or not buyer.last_name))):
        warnings.append("В карточке покупателя МойСклад не заполнены ИНН/КПП, юридический адрес или ФИО ИП. Заполните реквизиты в МойСклад и откройте УПД заново.")
    if source.get("invoice_unavailable"):
        warnings.append("Нет доступа к связанному счёту-фактуре. Обновите дескриптор и переустановите решение МойСклад для получения новых прав.")
    return {"seller": seller.model_dump(), "buyer": buyer.model_dump(),
            "number": invoice.get("name") or ("" if source.get("invoice_unavailable") else source.get("name", "")),
            "document_date": str(invoice.get("moment") or source.get("moment", ""))[:10],
            "function": "СЧФДОП" if source.get("factureOut") else "ДОП",
            "invoice_locked": bool(invoice), "warnings": warnings,
            "shipment_name": source.get("name"), "total": source.get("sum"),
            "positions": (source.get("positions") or {}).get("meta", {}).get("size", 0)}


@router.post("/{document_id}/export.upd.xml")
async def export_document_upd(document_id: UUID, options: UpdExportOptions,
                              current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    from urllib.parse import quote
    from starlette.concurrency import run_in_threadpool
    from app.services.upd_export import generate_upd, UpdExportError
    source = await _upd_source(document_id, current_user, db)
    try:
        filename, xml = await run_in_threadpool(generate_upd, source, options)
    except UpdExportError as exc:
        raise HTTPException(422, str(exc)) from None
    except ValueError:
        raise HTTPException(422, "Реквизиты содержат недопустимые для XML символы") from None
    return Response(xml, media_type="application/xml", headers={
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
        "Cache-Control": "no-store"})


@router.get("/{document_id}/upd-draft")
async def outgoing_upd_status(document_id: UUID, current_user: User = Depends(get_current_user),
                              db: AsyncSession = Depends(get_db), response: Response = None):
    from app.db.models import OutgoingUpd
    from app.services.outgoing_upd import response as draft_response
    doc = await db.scalar(select(Document.id).where(Document.id == document_id, Document.user_id == current_user.id))
    if not doc:
        raise HTTPException(404, "Документ не найден")
    row = await db.scalar(select(OutgoingUpd).where(OutgoingUpd.document_id == document_id, OutgoingUpd.user_id == current_user.id))
    if response is not None:
        response.headers['Cache-Control'] = 'no-store'
    return draft_response(row) if row else None


@router.post("/{document_id}/upd-draft")
async def create_outgoing_upd_draft(document_id: UUID, options: UpdExportOptions,
                                   current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    import hashlib
    from uuid import uuid4
    from app.db.models import OutgoingUpd
    from app.services.upd_export import generate_upd, UpdExportError
    from app.services.outgoing_upd import deliver, connection_key, response as draft_response
    from app.api.mark_control import _client_from_integration, _get_auth
    from app.services.saby import SabyError
    # Serialize first creation per shipment, including a repeated POST on another
    # backend process. A saved UUID survives crashes and lost HTTP responses.
    doc = await db.scalar(select(Document).where(Document.id == document_id, Document.user_id == current_user.id).with_for_update())
    if not doc:
        raise HTTPException(404, 'Документ не найден')
    if doc.kind != DocumentKind.demand or not doc.moysklad_id or doc.status != DocumentStatus.accepted:
        raise HTTPException(409, 'Сначала завершите запись марок отгрузки в МойСклад')
    row = await db.scalar(select(OutgoingUpd).where(OutgoingUpd.document_id == document_id))
    if row and row.status == 'created':
        return draft_response(row)
    integration = await db.scalar(select(Integration).where(Integration.user_id == current_user.id))
    client = _client_from_integration(integration)
    binding = connection_key(integration)
    if row and row.connection_key != binding:
        raise HTTPException(409, 'Подключение СБИС изменилось. Сверьте ранее созданный документ в исходном аккаунте')
    try:
        auth = await _get_auth(current_user.id, client)
    except (SabyError, httpx.HTTPError):
        raise HTTPException(502, 'Не удалось войти в СБИС. Проверьте подключение в разделе «Контроль марок»') from None
    replacement = None
    if row is None or row.status == 'failed':
        source = await _upd_source(document_id, current_user, db)
        if row is None:
            identifier = uuid4()
            try:
                from starlette.concurrency import run_in_threadpool
                filename, xml = await run_in_threadpool(generate_upd, source, options, file_uuid=identifier)
            except UpdExportError as exc:
                raise HTTPException(422, str(exc)) from None
            except ValueError:
                raise HTTPException(422, 'Реквизиты содержат недопустимые для XML символы') from None
            row = OutgoingUpd(id=identifier, document_id=document_id, user_id=current_user.id,
                              status='prepared', connection_key=binding, options=options.model_dump(mode='json'), filename=filename,
                              xml_text=xml.decode('cp1251'), xml_sha256=hashlib.sha256(xml).hexdigest())
            db.add(row)
        else:
            replacement = (source, options)
    await db.commit()
    row = await deliver(db, row.id, client, auth, replacement=replacement)
    return draft_response(row)


@router.get("/{document_id}/upd-draft.xml")
async def download_outgoing_upd_xml(document_id: UUID, current_user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    from urllib.parse import quote
    from app.db.models import OutgoingUpd
    row = await db.scalar(select(OutgoingUpd).where(OutgoingUpd.document_id == document_id, OutgoingUpd.user_id == current_user.id))
    if not row:
        raise HTTPException(404, 'Сохранённый XML УПД не найден')
    return Response(row.xml_text.encode('cp1251'), media_type='application/xml', headers={
        'Content-Disposition': f"attachment; filename*=UTF-8''{quote(row.filename)}", 'Cache-Control': 'no-store'})


@router.get("/{document_id}/export.xlsx")
async def export_document_xlsx(
    document_id: UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Выгрузить структуру документа в XLSX: наименование позиции + марка + кол-во.

    Агрегаты (короб/блок) разворачиваются в марки пачек (child_codes). Немаркированный
    товар (is_barcode) — строка с пустой маркой и количеством. SSCC «целиком» (is_box без
    child_codes) — строка с кодом SSCC и количеством. Включаются все сканы документа.
    """
    import io
    from urllib.parse import quote
    from fastapi.responses import Response
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    # Коды маркировки содержат управляющий разделитель GS (\x1d) и пр. control-символы —
    # openpyxl их запрещает (IllegalCharacterError). Вырезаем перед записью в ячейку.
    def _xl(v):
        return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

    result = await db.execute(
        select(Document).where(
            Document.id == document_id,
            Document.user_id == current_user.id,
        )
    )
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    scans_res = await db.execute(
        select(Scan).where(Scan.document_id == document_id).order_by(Scan.scanned_at)
    )
    scans = scans_res.scalars().all()

    rows: List[tuple] = []
    for s in scans:
        name = s.product_name or "—"
        if s.child_codes and not (s.is_box or s.keep_aggregate):
            for cc in s.child_codes:
                rows.append((name, cc, 1))
        elif s.is_barcode:  # немаркированный товар — марка пустая, кол-во из скана
            rows.append((name, "", int(s.box_quantity or 1)))
        elif s.is_box:  # SSCC «целиком» — код короба + кол-во внутри
            rows.append((name, s.code, int(s.box_quantity or 1)))
        else:
            rows.append((name, s.code, int(s.box_quantity or 1)))
    rows.sort(key=lambda r: (r[0], r[1]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Отгрузка"
    ws.append(["Наименование позиции", "Марка", "Кол-во"])
    for r in rows:
        ws.append([_xl(r[0]), _xl(r[1]), r[2]])
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 8
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)

    safe_name = (doc.name or "Документ").replace("/", "-").replace("\\", "-")
    filename = f"{safe_name}.xlsx"
    headers = {
        "Content-Disposition": (
            f"attachment; filename=order.xlsx; filename*=UTF-8''{quote(filename)}"
        )
    }
    return Response(
        content=buf.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers=headers,
    )


@router.post('/{document_id}/writeoff', response_model=DocumentResponse)
async def open_shipment_writeoff(document_id: UUID, current_user: User = Depends(get_current_user),
                                 db: AsyncSession = Depends(get_db)):
    """Open one local writeoff draft and copy marks from this exact completed shipment."""
    import copy
    from app.services.document_guard import lock_ms_document
    from app.services.writeoff_import import import_shipment_codes as import_codes
    source = (await db.execute(select(Document).where(
        Document.id == document_id, Document.user_id == current_user.id))).scalar_one_or_none()
    if source is None:
        raise HTTPException(404, 'Отгрузка не найдена')
    if source.kind != DocumentKind.demand or not source.moysklad_id:
        raise HTTPException(400, 'Выберите отгрузку, связанную с МойСкладом.')
    await lock_ms_document(db, current_user.id, DocumentKind.loss, source.moysklad_id)
    target = (await db.execute(select(Document).where(
        Document.user_id == current_user.id, Document.kind == DocumentKind.loss,
        Document.organization_profile_id == source.organization_profile_id,
        Document.moysklad_id == source.moysklad_id, Document.status != DocumentStatus.accepted,
        Document.upd_meta['superseded_by_document_id'].astext.is_(None),
    ).order_by(Document.created_at.desc()).limit(1).with_for_update())).scalar_one_or_none()
    if target is None:
        # Freeze the plan snapshot against a simultaneous refresh before copying it.
        source = (await db.execute(select(Document).where(
            Document.id == document_id, Document.user_id == current_user.id)
            .with_for_update().execution_options(populate_existing=True))).scalar_one()
        target = Document(user_id=current_user.id, kind=DocumentKind.loss, name=source.name,
            moysklad_id=source.moysklad_id, organization_profile_id=source.organization_profile_id,
            moysklad_organization_id=source.moysklad_organization_id, moysklad_store_id=source.moysklad_store_id,
            plan=copy.deepcopy(source.plan or []), upd_meta={'source_shipment_document_id': str(source.id)})
        db.add(target)
        await db.flush()
    try:
        await import_codes(db, target.id, current_user.id, source_id=source.id)
    except Exception:
        # Do not leave an empty draft if the source is unfinished or import is blocked.
        await db.rollback()
        raise
    await db.refresh(target)
    from app.services.scan_events import publish_event
    await publish_event(current_user.id, {'type': 'scans_changed', 'document_id': str(target.id)})
    return _doc_to_response(target, await _scan_count(db, target.id))


@router.post('/{document_id}/import-shipment-codes')
async def import_shipment_codes(document_id: UUID, current_user: User = Depends(get_current_user),
                                db: AsyncSession = Depends(get_db)):
    from app.services.writeoff_import import import_shipment_codes as import_codes
    result = await import_codes(db, document_id, current_user.id)
    from app.services.scan_events import publish_event
    await publish_event(current_user.id, {'type': 'scans_changed', 'document_id': str(document_id)})
    return result


@router.post("/{document_id}/verify")
async def verify_document(
    document_id: UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
    recheck_all: bool = False,
):
    """Пакетно проверить в ЧЗ все локально отсканированные марки (status=scanned).

    Основной флоу: скан ставит `scanned` (только формат GS1, без ЧЗ), а проверка
    в ЧЗ идёт здесь одной операцией. Прогресс — через WS scan_update по каждому скану,
    завершение — событие verify_done.
    """
    from app.services.document_guard import editable_document
    doc = await editable_document(db, document_id, current_user.id)
    if getattr(doc, 'kind', None) == DocumentKind.supply:
        from app.services.acceptance_mode import acceptance_cz_check_enabled
        if not await acceptance_cz_check_enabled(db, doc):
            raise HTTPException(409, 'Проверка ЧЗ для приёмки отключена в настройках юрлица.')

    from sqlalchemy import func
    from app.db.models import Scan, ScanStatus

    from app.services.shipment_corrections import correction
    baseline_ids = list((correction(doc) or {}).get('baseline_scans', {}))
    filters = [Scan.document_id == document_id, Scan.is_barcode.is_(False),
               Scan.status.in_([ScanStatus.scanned, ScanStatus.valid, ScanStatus.overflow, ScanStatus.invalid, ScanStatus.unknown_product])
               if recheck_all else Scan.status == ScanStatus.scanned]
    if baseline_ids:
        filters.append(Scan.id.not_in([UUID(value) for value in baseline_ids]))
    pending_q = await db.execute(
        select(func.count(Scan.id)).where(
            *filters,
        )
    )
    to_check = pending_q.scalar_one()

    if recheck_all:
        selected = (await db.execute(select(Scan).where(*filters))).scalars().all()
        for scan in selected:
            scan.status = ScanStatus.scanned
            scan.verified_at = None
            scan.verification = {'source': 'pending', 'checked_at': None,
                                 'owner_result': 'unknown', 'owner_reason': 'Идёт повторная проверка ЧЗ'}
        await db.commit()
        from app.services.scan_events import publish_event
        await publish_event(current_user.id, {'type': 'scans_changed', 'document_id': str(doc.id)})

    from app.worker.tasks import verify_document_task
    verify_document_task.delay(str(document_id), str(current_user.id), recheck_all)
    return {"status": "verifying", "document_id": str(document_id), "count": to_check}


@router.post("/{document_id}/process")
async def process_document(
    document_id: UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """
    Завершить документ: Celery обновляет МС (positions, при необходимости
    trackingCodes) и ставит статус accepted. API Честного Знака не вызывается.
    """
    from app.services.document_guard import editable_document
    doc = await editable_document(db, document_id, current_user.id)
    if doc.kind not in (DocumentKind.demand, DocumentKind.supply):
        raise HTTPException(400, "Для списания используйте отправку в Честный Знак")
    if (getattr(doc, 'upd_meta', None) or {}).get('shipment_correction'):
        raise HTTPException(409, 'Сохраните изменения через раздел «Исправить отгрузку». Полная отправка исправления запрещена.')
    require_collection_started(doc)
    skip_cz = False
    if doc.kind == DocumentKind.supply:
        from app.services.acceptance_mode import acceptance_cz_check_enabled
        skip_cz = not await acceptance_cz_check_enabled(db, doc)
        if skip_cz:
            from app.services.supply_positions import prepare_supply_without_cz
            try:
                prepare_supply_without_cz(doc.plan, getattr(doc, 'upd_meta', None))
            except ValueError as exc:
                raise HTTPException(409, str(exc)) from exc
    if not doc.moysklad_id:
        raise HTTPException(409, "Сначала выберите документ МойСклад")
    ms = await _get_ms_service(current_user, db)
    if doc.kind == DocumentKind.demand:
        from app.services.shipment_guard import ensure_active_shipment
        try:
            shipment = await ms.get_document('demand', doc.moysklad_id)
        except Exception as exc:
            raise HTTPException(502, 'Не удалось проверить отгрузку в МойСкладе. Повторите попытку.') from exc
        try:
            ensure_active_shipment(shipment)
        except ValueError as exc:
            raise HTTPException(409, str(exc)) from exc

    from sqlalchemy import func
    from app.db.models import Scan, ScanStatus

    # Не отправляем в МС непроверенные марки: сначала «Проверить марки».
    scanned_q = await db.execute(
        select(func.count(Scan.id)).where(
            Scan.document_id == document_id,
            Scan.status.in_([ScanStatus.scanned, ScanStatus.pending]),
        )
    )
    scanned_count = scanned_q.scalar_one()
    if scanned_count and not skip_cz:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Сначала проверьте марки ({scanned_count} не проверено) — "
                f"нажмите «Проверить марки»"
            ),
        )

    # Блокируем процесс, если есть сканы со статусом unknown_product —
    # их нельзя отправить в МС без сопоставления товара.
    unknown_q = await db.execute(
        select(func.count(Scan.id)).where(
            Scan.document_id == document_id,
            Scan.status == ScanStatus.unknown_product,
        )
    )
    unknown_count = unknown_q.scalar_one()
    if unknown_count and not skip_cz:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Сначала сопоставьте товары для {unknown_count} "
                f"кодов с неизвестным GTIN"
            ),
        )

    from app.worker.tasks import process_document_task
    count = (await db.execute(select(func.count(Scan.id)).where(
        Scan.document_id == document_id, Scan.status.in_([ScanStatus.valid, ScanStatus.overflow])
    ))).scalar_one()
    from app.services.supply_positions import prepare_supply_rows
    quantity_only = prepare_supply_rows([], doc.plan, getattr(doc, 'upd_meta', None), include_marked=skip_cz) if doc.kind == DocumentKind.supply else []
    if not count and not quantity_only:
        raise HTTPException(409, "Нет проверенных марок или товаров для отправки")
    doc.status = DocumentStatus.processing
    if doc.kind == DocumentKind.supply:
        doc.upd_meta = {**(doc.upd_meta or {}), 'acceptance_cz_check_enabled': not skip_cz}
    doc.error_message = None
    doc.processing_progress = {"sent": 0, "total": count, "stage": "preparing"}
    await db.commit()

    try:
        process_document_task.delay(str(document_id), str(current_user.id))
    except Exception:
        doc.status = DocumentStatus.draft
        doc.error_message = "Не удалось поставить отправку в очередь. Повторите попытку."
        await db.commit()
        raise HTTPException(503, doc.error_message)
    return {"status": "processing", "document_id": str(document_id)}


@router.post("/{document_id}/accept")
async def accept_document_alias(
    document_id: UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Backward-совместимый алиас → /process."""
    return await process_document(document_id, current_user, db)


@router.post('/{document_id}/start-collection')
async def begin_collection(document_id: UUID, current_user: User = Depends(get_current_user),
                           db: AsyncSession = Depends(get_db)):
    doc = (await db.execute(select(Document).where(
        Document.id == document_id, Document.user_id == current_user.id,
    ))).scalar_one_or_none()
    if not doc:
        raise HTTPException(404, 'Документ не найден')
    if doc.kind != DocumentKind.demand or not doc.moysklad_id:
        return {'status': 'started'}
    from app.services.collection_start import start_collection
    await start_collection(db, doc, await _get_ms_service(current_user, db))
    return {'status': 'started', 'collection_started': True}
