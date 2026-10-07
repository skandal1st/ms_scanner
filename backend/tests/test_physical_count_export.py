import io
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services.physical_count_export import build_count_xlsx, report_rows


def session(method='scan', reviewed=None):
    return SimpleNamespace(mode='inventory', status='active', name='Сверка склада',
        settings={'count_method': method, 'reviewed_brands': reviewed or [], 'store_name': 'Основной'},
        plan=[{'key': 'p', 'folder_name': 'Бренд', 'product_name': '=HYPERLINK("bad")\x1d',
               'gtins': ['04620543080527'], 'base_qty': 10, 'shipment_qty': 2, 'expected_qty': 12},
              {'key': 'empty', 'folder_name': 'Другой', 'product_name': 'Товар',
               'gtin': '00000000000001', 'base_qty': 3, 'shipment_qty': 0, 'expected_qty': 3}])


def test_xlsx_preserves_gtin_numbers_and_unreviewed_difference():
    workbook = load_workbook(io.BytesIO(build_count_xlsx(session(), {'p': 9})))
    sheet = workbook.active
    assert sheet['B7'].value == '=HYPERLINK("bad")'
    assert sheet['B7'].data_type == 's'
    assert sheet['C7'].value == '04620543080527'
    assert sheet['C8'].value == '00000000000001'
    assert [sheet.cell(7, col).value for col in range(4, 10)] == [10, 2, 12, 9, None, 'Нет']
    assert sheet['G8'].value == 0 and sheet['H8'].value is None
    assert sheet.freeze_panes == 'D7' and sheet.auto_filter.ref == 'A6:I8'
    workbook.close()


def test_reviewed_scan_and_explicit_zero_are_distinct_from_unentered():
    rows = list(report_rows(session(reviewed=['Бренд']), {'p': 9}))
    assert rows[0][6:] == [9, -3, 'Да']
    manual = session('quantity')
    rows = list(report_rows(manual, {'p': 0}))
    assert rows[0][6:] == [0, -12, 'Да']
    assert rows[1][6:] == [None, None, 'Нет']
    workbook = load_workbook(io.BytesIO(build_count_xlsx(manual, {'p': 0})))
    assert workbook.active['G7'].value == 0
    assert workbook.active['G8'].value is None
    workbook.close()
